#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BraTS 2023 Missing-Modality: Full-only vs Modality Dropout
Generate reproducible CSV, Excel and paper-style figures.

Run:
  python generate_moddrop_paper_assets.py
"""
from __future__ import annotations

import json, math
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import ColorScaleRule
    HAVE_OPENPYXL = True
except Exception:
    HAVE_OPENPYXL = False

# ============================== CONFIG ==============================
PROJECT_ROOT = Path("/root/autodl-tmp/nnunet_brats")
BENCHMARK_ROOT = PROJECT_ROOT / "missing_modality_benchmark"
OUTPUT_DIR = BENCHMARK_ROOT / "paper_assets_moddrop"
DPI = 300
REGIONS = ["WT", "TC", "ET"]
MODELS = ["Full-only", "Modality Dropout"]
SETTINGS = ["Full input", "Missing T1c", "Missing FLAIR", "Missing T2w", "Missing T1n"]
MISSING = SETTINGS[1:]

# Used only when a real JSON cannot be located or parsed.
FALLBACK = {
"Full-only": {
"Full input":{"WT":(0.9326,0.0839),"TC":(0.9106,0.1645),"ET":(0.8669,0.1943),"Mean":(0.9034,np.nan)},
"Missing T1c":{"WT":(0.9019,0.1116),"TC":(0.2648,0.2955),"ET":(0.0599,0.1643),"Mean":(0.4089,np.nan)},
"Missing FLAIR":{"WT":(0.7070,0.2335),"TC":(0.8625,0.2317),"ET":(0.8308,0.2296),"Mean":(0.8001,np.nan)},
"Missing T2w":{"WT":(0.8334,0.2183),"TC":(0.8578,0.2462),"ET":(0.8391,0.2106),"Mean":(0.8434,np.nan)},
"Missing T1n":{"WT":(0.9287,0.0613),"TC":(0.9057,0.1547),"ET":(0.8585,0.1837),"Mean":(0.8976,np.nan)},
},
"Modality Dropout": {
"Full input":{"WT":(0.9360,0.0835),"TC":(0.9133,0.1558),"ET":(0.8723,0.1833),"Mean":(0.9072,np.nan)},
"Missing T1c":{"WT":(0.9272,0.0887),"TC":(0.8095,0.1782),"ET":(0.6540,0.2179),"Mean":(0.7969,np.nan)},
"Missing FLAIR":{"WT":(0.9016,0.1027),"TC":(0.9024,0.1763),"ET":(0.8626,0.1956),"Mean":(0.8889,np.nan)},
"Missing T2w":{"WT":(0.9290,0.0855),"TC":(0.9049,0.1689),"ET":(0.8619,0.1967),"Mean":(0.8986,np.nan)},
"Missing T1n":{"WT":(0.9338,0.0839),"TC":(0.9059,0.1657),"ET":(0.8660,0.1868),"Mean":(0.9019,np.nan)},
}}

CANDIDATES = {
("Full-only","Full input"):[BENCHMARK_ROOT/"fold_0_full_best"/"metrics",BENCHMARK_ROOT/"fold_0_full_best"],
("Full-only","Missing T1c"):[BENCHMARK_ROOT/"fold_0_missing_T1c"/"metrics",BENCHMARK_ROOT/"fold_0_missing_T1c"],
("Full-only","Missing FLAIR"):[BENCHMARK_ROOT/"fold_0_missing_FLAIR"/"metrics",BENCHMARK_ROOT/"fold_0_missing_FLAIR"],
("Full-only","Missing T2w"):[BENCHMARK_ROOT/"fold_0_missing_T2w"/"metrics",BENCHMARK_ROOT/"fold_0_missing_T2w"],
("Full-only","Missing T1n"):[BENCHMARK_ROOT/"fold_0_missing_T1n"/"metrics",BENCHMARK_ROOT/"fold_0_missing_T1n"],
("Modality Dropout","Full input"):[
PROJECT_ROOT/"nnUNet_results"/"Dataset501_BraTS2023"/"nnUNetTrainerModalityDropout__nnUNetPlans__3d_fullres"/"fold_0"/"metrics",
PROJECT_ROOT/"nnUNet_results"/"Dataset501_BraTS2023"/"nnUNetTrainerModalityDropout__nnUNetPlans__3d_fullres"/"fold_0"],
("Modality Dropout","Missing T1c"):[BENCHMARK_ROOT/"moddrop_fold_0_missing_T1c"/"metrics",BENCHMARK_ROOT/"moddrop_fold_0_missing_T1c"],
("Modality Dropout","Missing FLAIR"):[BENCHMARK_ROOT/"moddrop_fold_0_missing_FLAIR"/"metrics",BENCHMARK_ROOT/"moddrop_fold_0_missing_FLAIR"],
("Modality Dropout","Missing T2w"):[BENCHMARK_ROOT/"moddrop_fold_0_missing_T2w"/"metrics",BENCHMARK_ROOT/"moddrop_fold_0_missing_T2w"],
("Modality Dropout","Missing T1n"):[BENCHMARK_ROOT/"moddrop_fold_0_missing_T1n"/"metrics",BENCHMARK_ROOT/"moddrop_fold_0_missing_T1n"],
}
# =====================================================================

LOG: list[str] = []
def log(s: str, level="INFO"):
    line = f"[{level}] {s}"
    LOG.append(line); print(line)
def finite(x): 
    try: return x is not None and math.isfinite(float(x))
    except Exception: return False
def f4(x): return f"{float(x):.4f}" if finite(x) else "NA"
def nk(s): return str(s).replace("_"," ").replace("-"," ").strip().lower()

def find_value(obj: Any, names: list[str]):
    wanted={nk(n) for n in names}
    if isinstance(obj,dict):
        for k,v in obj.items():
            if nk(k) in wanted: return v
        for v in obj.values():
            ret=find_value(v,names)
            if ret is not None: return ret
    return None

def meanstd(v):
    if isinstance(v,(int,float)): return float(v),np.nan
    if not isinstance(v,dict): return np.nan,np.nan
    mean=np.nan; std=np.nan
    for k,x in v.items():
        if nk(k) in {"mean","avg","average","dice mean","score","value"}: mean=float(x)
        if nk(k) in {"std","stdev","standard deviation","dice std"}: std=float(x)
    if finite(mean): return mean,std
    if len(v)==1: return meanstd(next(iter(v.values())))
    return np.nan,np.nan

def aliases(region):
    return {
    "WT":["WT","wt","whole tumor","whole_tumor"],
    "TC":["TC","tc","tumor core","tumor_core"],
    "ET":["ET","et","enhancing tumor","enhancing_tumor"],
    "Mean":["Mean","mean","mean dice","mean_dice"],
    }[region]

def parse_json(path: Path):
    try: d=json.loads(path.read_text(encoding="utf-8"))
    except Exception as e: log(f"Cannot parse {path}: {e}","WARNING"); return None
    root=find_value(d,["all_cases","all cases"])
    if not isinstance(root,dict): root=d
    out={}
    for r in ["WT","TC","ET","Mean"]:
        m,s=meanstd(find_value(root,aliases(r)))
        if not finite(m) and isinstance(root,dict):
            # flat keys e.g. wt_mean
            keys={nk(x).replace(" ","_") for x in aliases(r)}
            for k,v in root.items():
                q=nk(k).replace(" ","_")
                if q in {z+"_mean" for z in keys}: m=float(v)
                if q in {z+"_std" for z in keys}: s=float(v)
        out[r]=(m,s)
    return out if sum(finite(out[r][0]) for r in out)>=3 else None

def metric_files(dirs):
    found=[]; seen=set()
    for d in dirs:
        if not d.exists(): continue
        if d.is_file() and d.suffix==".json": found.append(d); continue
        for p in [d/"wt_tc_et_summary.json",d/"summary.json",d/"metrics.json"]:
            if p.exists() and p not in seen: found.append(p);seen.add(p)
        for child in [d/"metrics",d/"validation"]:
            if child.exists():
                for p in child.glob("*.json"):
                    if p not in seen: found.append(p);seen.add(p)
    return found

def load_metrics(model,setting):
    for p in metric_files(CANDIDATES[(model,setting)]):
        parsed=parse_json(p)
        if parsed:
            log(f"{model} | {setting}: loaded {p}")
            for r in ["WT","TC","ET","Mean"]:
                if finite(parsed[r][0]) and abs(parsed[r][0]-FALLBACK[model][setting][r][0])>1e-4:
                    log(f"{model} | {setting} {r}: JSON={parsed[r][0]:.4f}, verified fallback={FALLBACK[model][setting][r][0]:.4f}; JSON retained.","WARNING")
            return parsed,str(p),False
    log(f"{model} | {setting}: fallback values used because metric JSON unavailable/unparseable.","WARNING")
    return FALLBACK[model][setting],"FALLBACK",True

def build_frames():
    rows=[]
    for model in MODELS:
        for setting in SETTINGS:
            vals,src,fb=load_metrics(model,setting)
            for r in ["WT","TC","ET","Mean"]:
                m,s=vals[r]
                rows.append(dict(model=model,setting=setting,missing_modality="None" if setting=="Full input" else setting.replace("Missing ",""),
                                 metric_group="all_cases",region=r,dice_mean=m,dice_std=s,source_json=src,used_fallback=fb))
    long=pd.DataFrame(rows)
    rows=[]
    for model in MODELS:
        ref=long[(long.model==model)&(long.setting=="Full input")].set_index("region")["dice_mean"].to_dict()
        for setting in SETTINGS:
            q=long[(long.model==model)&(long.setting==setting)].set_index("region")
            row={"Model":model,"Input Setting":setting,"Missing Modality":"None" if setting=="Full input" else setting.replace("Missing ","")}
            for r in ["WT","TC","ET","Mean"]:
                m=q.loc[r,"dice_mean"]; s=q.loc[r,"dice_std"]
                row[f"{r}_mean"]=m; row[f"{r}_std"]=s
                row[f"{r}_display"]=f"{m:.4f} ± {s:.4f}" if finite(s) else f"{m:.4f}"
                row[f"{r}_drop_from_model_full"]=0.0 if setting=="Full input" else ref[r]-m
            rows.append(row)
    wide=pd.DataFrame(rows)
    f=wide[wide.Model=="Full-only"].set_index("Input Setting")
    m=wide[wide.Model=="Modality Dropout"].set_index("Input Setting")
    absrows=[]; droprows=[]; recr=[]
    for st in SETTINGS:
        z={"Input Setting":st}
        for r in ["WT","TC","ET","Mean"]:
            z[f"Full-only {r}"]=f.loc[st,f"{r}_mean"];z[f"ModDrop {r}"]=m.loc[st,f"{r}_mean"];z[f"ModDrop - Full-only {r}"]=m.loc[st,f"{r}_mean"]-f.loc[st,f"{r}_mean"]
        absrows.append(z)
    for st in MISSING:
        z={"Missing Modality":st.replace("Missing ","")}; rec={"Missing Modality":st.replace("Missing ","")}
        for r in ["WT","TC","ET","Mean"]:
            fd=f.loc[st,f"{r}_drop_from_model_full"]; md=m.loc[st,f"{r}_drop_from_model_full"]
            z[f"Full-only {r} Drop"]=fd;z[f"ModDrop {r} Drop"]=md;z[f"Drop Reduction {r}"]=fd-md
            rec[f"{r} Recovery Gain"]=m.loc[st,f"{r}_mean"]-f.loc[st,f"{r}_mean"]
        droprows.append(z);recr.append(rec)
    return long,wide,pd.DataFrame(absrows),pd.DataFrame(droprows),pd.DataFrame(recr)

def annotate(ax,bars):
    for b in bars:
        h=b.get_height()
        ax.annotate(f"{h:.3f}",(b.get_x()+b.get_width()/2,h),xytext=(0,3),textcoords="offset points",ha="center",va="bottom",fontsize=7)

def plot_absolute(wide):
    f=wide[wide.Model=="Full-only"].set_index("Input Setting");m=wide[wide.Model=="Modality Dropout"].set_index("Input Setting")
    x=np.arange(len(SETTINGS));w=.36;fig,axs=plt.subplots(2,2,figsize=(12,7),sharex=True)
    for ax,r in zip(axs.flat,["WT","TC","ET","Mean"]):
        a=[f.loc[s,f"{r}_mean"] for s in SETTINGS];b=[m.loc[s,f"{r}_mean"] for s in SETTINGS]
        ea=None if r=="Mean" else [f.loc[s,f"{r}_std"] for s in SETTINGS];eb=None if r=="Mean" else [m.loc[s,f"{r}_std"] for s in SETTINGS]
        p=ax.bar(x-w/2,a,w,label="Full-only",yerr=ea,capsize=2);q=ax.bar(x+w/2,b,w,label="Modality Dropout",yerr=eb,capsize=2);annotate(ax,p);annotate(ax,q)
        ax.set_title(f"{r} Dice");ax.set_ylim(0,1.05);ax.set_ylabel("Dice");ax.set_xticks(x,SETTINGS,rotation=20,ha="right");ax.grid(axis="y",alpha=.25)
        if r=="WT":ax.legend(loc="lower left")
    fig.suptitle("Absolute Dice Comparison",y=1.02);fig.tight_layout();p=OUTPUT_DIR/"absolute_dice_fullonly_vs_moddrop.png";fig.savefig(p,dpi=DPI,bbox_inches="tight");plt.close(fig);return p

def plot_drop(drop):
    x=np.arange(len(MISSING));w=.36;fig,axs=plt.subplots(2,2,figsize=(12,7),sharex=True)
    for ax,r in zip(axs.flat,["WT","TC","ET","Mean"]):
        a=drop[f"Full-only {r} Drop"];b=drop[f"ModDrop {r} Drop"]
        p=ax.bar(x-w/2,a,w,label="Full-only");q=ax.bar(x+w/2,b,w,label="Modality Dropout");annotate(ax,p);annotate(ax,q)
        ax.set_title(f"{r} Dice Drop");ax.set_ylabel("Dice Drop from Complete-Modality Input");ax.set_xticks(x,[s.replace("Missing ","") for s in MISSING],rotation=20,ha="right");ax.grid(axis="y",alpha=.25)
        if r=="WT":ax.legend()
    fig.suptitle("Robustness Dice Drop Comparison",y=1.02);fig.tight_layout();p=OUTPUT_DIR/"dice_drop_fullonly_vs_moddrop.png";fig.savefig(p,dpi=DPI,bbox_inches="tight");plt.close(fig);return p

def heat(mat,rows,title,path,vmin,vmax,label):
    fig,ax=plt.subplots(figsize=(7,4.6));im=ax.imshow(mat,aspect="auto",vmin=vmin,vmax=vmax);cb=fig.colorbar(im,ax=ax);cb.set_label(label)
    ax.set_xticks(np.arange(3),REGIONS);ax.set_yticks(np.arange(len(rows)),rows);ax.set_xlabel("Tumor Region");ax.set_ylabel("Missing Modality");ax.set_title(title)
    for i in range(mat.shape[0]):
        for j in range(mat.shape[1]):ax.text(j,i,f"{mat[i,j]:.4f}",ha="center",va="center",fontsize=9)
    fig.tight_layout();fig.savefig(path,dpi=DPI,bbox_inches="tight");plt.close(fig);return path

def plot_heatmaps(drop,recovery):
    rows=drop["Missing Modality"].tolist();fm=drop[[f"Full-only {r} Drop" for r in REGIONS]].to_numpy(float);mm=drop[[f"ModDrop {r} Drop" for r in REGIONS]].to_numpy(float);mx=max(np.nanmax(fm),np.nanmax(mm),.01)
    a=heat(fm,rows,"Modality-Region Risk Matrix: Full-only Model",OUTPUT_DIR/"modality_region_risk_heatmap_fullonly.png",0,mx,"Dice Drop")
    b=heat(mm,rows,"Modality-Region Risk Matrix: Modality Dropout Model",OUTPUT_DIR/"modality_region_risk_heatmap_moddrop.png",0,mx,"Dice Drop")
    rm=recovery[[f"{r} Recovery Gain" for r in REGIONS]].to_numpy(float);lim=max(abs(rm.min()),abs(rm.max()),.01)
    c=heat(rm,rows,"Recovery Gain of Modality Dropout over Full-only",OUTPUT_DIR/"modality_dropout_recovery_heatmap.png",-lim,lim,"Recovery Gain")
    return [a,b,c]

def plot_mean(wide):
    f=wide[wide.Model=="Full-only"].set_index("Input Setting");m=wide[wide.Model=="Modality Dropout"].set_index("Input Setting");x=np.arange(len(SETTINGS));w=.36;fig,ax=plt.subplots(figsize=(9,5))
    p=ax.bar(x-w/2,[f.loc[s,"Mean_mean"] for s in SETTINGS],w,label="Full-only");q=ax.bar(x+w/2,[m.loc[s,"Mean_mean"] for s in SETTINGS],w,label="Modality Dropout");annotate(ax,p);annotate(ax,q)
    ax.set_xticks(x,SETTINGS,rotation=20,ha="right");ax.set_ylim(0,1.05);ax.set_ylabel("Mean Dice");ax.set_title("Mean Dice Across Input Conditions");ax.legend();ax.grid(axis="y",alpha=.25);fig.tight_layout();z=OUTPUT_DIR/"mean_dice_comparison.png";fig.savefig(z,dpi=DPI,bbox_inches="tight");plt.close(fig);return z

def plot_recovery_table(rec):
    fig,ax=plt.subplots(figsize=(8.8,2.5));ax.axis("off");d=rec.copy()
    for c in d.columns[1:]:d[c]=d[c].map(lambda x:f"{x:+.4f}")
    t=ax.table(cellText=d.values,colLabels=d.columns,cellLoc="center",loc="center");t.auto_set_font_size(False);t.set_fontsize(8.5);t.scale(1,1.6);ax.set_title("Recovery Gain of Modality Dropout over Full-only",pad=12);fig.tight_layout();z=OUTPUT_DIR/"modality_dropout_recovery_table.png";fig.savefig(z,dpi=DPI,bbox_inches="tight");plt.close(fig);return z

def protocol():
    rows=[]
    for model in MODELS:
        for st in SETTINGS:
            rows.append({"Model":model,"Trainer":"nnUNetTrainer" if model=="Full-only" else "nnUNetTrainerModalityDropout",
            "Training Input Policy":"Complete 4-modality input only" if model=="Full-only" else "50% complete; 50% one randomly zero-filled modality",
            "Inference Input":st,"Missing Modality":"None" if st=="Full input" else st.replace("Missing ",""),"Fold":0,"Checkpoint":"checkpoint_best.pth","Validation Cases":251,"Evaluation Regions":"WT / TC / ET","Notes":"single-fold preliminary benchmark; zero-filled missing-modality direct-inference protocol"})
    return pd.DataFrame(rows)

def export_excel(long,wide,absolute,drop,recovery):
    p=OUTPUT_DIR/"all_results_summary.xlsx"
    with pd.ExcelWriter(p,engine="openpyxl") as x:
        absolute.round(4).to_excel(x,sheet_name="Absolute_Dice",index=False);drop.round(4).to_excel(x,sheet_name="Dice_Drop",index=False);recovery.round(4).to_excel(x,sheet_name="Recovery_Gain",index=False);protocol().to_excel(x,sheet_name="Protocol",index=False);long.round(4).to_excel(x,sheet_name="Raw_Long_Format",index=False);wide.round(4).to_excel(x,sheet_name="Wide_All_Results",index=False);pd.DataFrame({"Log":LOG}).to_excel(x,sheet_name="Validation_Log",index=False)
    if HAVE_OPENPYXL:
        wb=load_workbook(p);fill=PatternFill("solid",fgColor="D9EAF7")
        for ws in wb.worksheets:
            ws.freeze_panes="A2"
            for c in ws[1]:c.font=Font(bold=True);c.fill=fill;c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            for i,col in enumerate(ws.columns,1):ws.column_dimensions[get_column_letter(i)].width=min(max(max(len(str(c.value or "")) for c in col)+2,12),42)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value,(float,np.floating)):c.number_format="0.0000"
        for name in ["Dice_Drop","Recovery_Gain"]:
            ws=wb[name];ws.conditional_formatting.add(f"B2:{get_column_letter(ws.max_column)}{ws.max_row}",ColorScaleRule(start_type="min",start_color="FFFFFF",mid_type="percentile",mid_value=50,mid_color="FFF2CC",end_type="max",end_color="F4CCCC"))
        wb.save(p)
    return p

def write_readme(absolute,drop,recovery):
    fallback="\n".join("- "+x for x in LOG if "fallback values used" in x) or "- No fallback metrics were required."
    text=f"""# BraTS Missing-Modality Paper Assets

## Reproduction
```bash
cd {OUTPUT_DIR}
python generate_moddrop_paper_assets.py
```

## Data provenance
The script first searches existing metric JSON files. Fallback values are used only when source metrics cannot be read.

{fallback}

## 中文结果总结
在当前 zero-filled missing-modality direct-inference stress test 下，Full-only 模型呈现明显的模态—区域特异性退化：T1c 缺失主要损害 TC 与 ET，FLAIR 缺失主要损害 WT。Modality Dropout 缓解了四种单模态缺失造成的退化，并且在完整输入下未观察到明显性能代价。Missing T1c 下 ET 仍有残余下降，这为后续模态—区域风险校准训练或伪标签策略提供动机。

当前仅为 fold 0 的 single-fold preliminary benchmark，不应解释为最终临床缺失模态性能。后续仍需要 missing-modality-aware training、低标注实验、五折交叉验证、HD95、统计检验和病例可视化。

## English Results-style paragraph
Under the zero-filled missing-modality direct-inference protocol, the Full-only model showed a region-dependent degradation pattern: missing T1c predominantly affected TC and ET, whereas missing FLAIR mainly impaired WT. In this single-fold preliminary benchmark, Modality Dropout substantially mitigated performance loss across the evaluated missing-modality settings without an observable loss under complete-modality inference. Nevertheless, residual degradation remained for ET under missing T1c. These results should be interpreted as stress-test findings and require further validation under missing-modality-aware training, multi-fold evaluation, and clinical-style protocols.

## Main Tables

### Absolute Dice
{absolute.to_markdown(index=False,floatfmt=".4f")}

### Robustness Dice Drop
{drop.to_markdown(index=False,floatfmt=".4f")}

### Recovery Gain
{recovery.to_markdown(index=False,floatfmt=".4f")}
"""
    p=OUTPUT_DIR/"README_paper_assets.md";p.write_text(text,encoding="utf-8");return p

def main():
    OUTPUT_DIR.mkdir(parents=True,exist_ok=True);log(f"Started {datetime.now().isoformat(timespec='seconds')}");log(f"Output directory: {OUTPUT_DIR}")
    long,wide,absolute,drop,recovery=build_frames()
    a=OUTPUT_DIR/"all_results_long.csv";b=OUTPUT_DIR/"all_results_wide.csv";long.round(4).to_csv(a,index=False);wide.round(4).to_csv(b,index=False);log(f"Saved {a}");log(f"Saved {b}")
    x=export_excel(long,wide,absolute,drop,recovery);log(f"Saved {x}")
    figs=[plot_absolute(wide),plot_drop(drop),*plot_heatmaps(drop,recovery),plot_mean(wide),plot_recovery_table(recovery)]
    for p in figs:log(f"Saved figure {p}")
    r=write_readme(absolute,drop,recovery);log(f"Saved {r}")
    v=OUTPUT_DIR/"validation_log.txt";v.write_text("\n".join(LOG)+"\n",encoding="utf-8")
    # basic checks
    for p in [a,b,x,*figs,r,v]:
        if not p.exists() or p.stat().st_size==0:log(f"Missing/empty output {p}","WARNING")
    v.write_text("\n".join(LOG)+"\n",encoding="utf-8")
    print("\nGenerated files:")
    for p in sorted(OUTPUT_DIR.iterdir()):print(p.resolve())

if __name__=="__main__":
    main()
